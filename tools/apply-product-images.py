# -*- coding: utf-8 -*-
"""نوشتن تصاویر پوشهٔ docs/تصاویر-محصولات در ستون «تصاویر» فایل اکسل کاتالوگ.

    python3 tools/apply-product-images.py [فایل.xlsx]

تطبیق بر پایهٔ «کد محصول» (SKU) انجام می‌شود، نه نام یا شناسه — کد یکتاست و
تغییر نمی‌کند. برای هر محصول، تصویر اول شاخص و بقیه گالری است؛ همان قراردادی که
افزونهٔ همگام‌سازی از ستون «تصاویر» می‌خواند. جداکنندهٔ نام‌ها ویرگول لاتین
است — هم افزونه و هم درون‌ریز خود ووکامرس همین را می‌فهمند.

در سلول فقط *نام فایل* نوشته می‌شود، نه نشانی کامل. افزونه نام فایل را در
کتابخانهٔ رسانه و پوشهٔ uploads می‌گردد، پس کافی است همین پوشه یک بار در
کتابخانهٔ رسانهٔ وردپرس بارگذاری شود.

اسکریپت بی‌خطر است و چند بار اجرا شدنش اشکالی ندارد: فقط ستون «تصاویر» سطرهایی
را می‌نویسد که در نقشهٔ پایین آمده‌اند و به بقیهٔ سلول‌ها دست نمی‌زند.
"""
import os
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMAGE_DIR = os.path.join(ROOT, "docs", "تصاویر-محصولات")
WORKBOOK = os.path.join(ROOT, "docs", "محصولات-سایت.xlsx")

# کد محصول → تصاویر (اولی شاخص، بقیه گالری).
MAPPING = {
    # کارتن پستی — هر سایز عکس اختصاصی خودش را دارد؛ کد روی عکس با سایز محصول
    # خوانده و تطبیق داده شده است (کد ۱ = ۱۵×۱۰×۱۰ … کد ۹ = ۵۵×۴۵×۳۵).
    "postalcartons-15-10-10": ["postal-carton-code-1.png"],
    "postalcartons-20-15-10": ["postal-carton-code-2.png"],
    "postalcartons-20-20-15": ["postal-carton-code-3.png"],
    "postalcartons-30-20-20": ["postal-carton-code-4.png"],
    "postalcartons-35-25-20": ["postal-carton-code-5.png"],
    "postalcartons-45-25-20": ["postal-carton-code-6.png"],
    "postalcartons-40-30-25": ["postal-carton-code-7.png"],
    "postalcartons-45-40-30": ["postal-carton-code-8.png"],
    "postalcartons-55-45-35": ["postal-carton-code-9.png"],

    # کارتن اسباب‌کشی — سه عکس ساده برای سه سایز. نسخهٔ چاپ‌دار عکس چاپ‌دار را
    # شاخص می‌گیرد و عکس سایز خودش را در گالری، تا اندازه هم دیده شود.
    "movingcartons-50-30-35": ["moving-carton-small.jpg"],
    "movingcartons-60-40-40": ["moving-carton-medium.jpg"],
    "movingcartons-70-50-40": ["moving-carton-large.jpg"],
    "movingcartons-50-30-35-print": ["moving-carton-printed.jpg", "moving-carton-small.jpg"],
    "movingcartons-60-40-40-print": ["moving-carton-printed.jpg", "moving-carton-medium.jpg"],
    "movingcartons-70-50-40-print": ["moving-carton-printed.jpg", "moving-carton-large.jpg"],

    # جعبهٔ پیتزا — عکس‌ها بر پایهٔ نامشان بازهٔ سایز ۲۰ تا ۴۲ را پوشش می‌دهند،
    # پس همهٔ این سایزها یک جفت عکس مشترک می‌گیرند.
    "ff-02": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-03": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-04": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-05": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-06": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-07": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-08": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-09": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-10": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-12": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-13": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    "ff-14": ["pizza-box-kraft-printed.png", "pizza-box-white-printed.png"],
    # «زیره و رویه» دو تکهٔ تخت است، پس عکس دستهٔ سینگل شاخص می‌شود.
    "ff-01": ["pizza-box-single.jpg", "pizza-box-kraft-printed.png"],
    "ff-11": ["pizza-box-single.jpg", "pizza-box-kraft-printed.png"],

    # جا لیوانی — عکس‌ها از روی نام فایلِ پوشه به محصول نسبت داده شده‌اند.
    "ju-01": ["cup-holder-2-kraft.png", "cup-holder-2-kraft-open.png"],
    "ju-02": ["cup-holder-2-printed.png", "cup-holder-2-printed-blue.jpg",
              "cup-holder-2-printed-colors.jpg"],
    "ju-03-vip": ["cup-holder-2-glued-vip.jpg"],
    "ju-06": ["cup-holder-4-kraft.png", "cup-holder-4-kraft-top.jpg",
              "cup-holder-4-kraft-in-use.png"],
    "ju-07": ["cup-holder-4-printed.png"],
    "ju-08-vip": ["cup-holder-4-glued-vip.jpg"],
    "ju-10": ["cup-holder-4-printed-cafe.png"],
    # ju-04 (دسته‌دار)، ju-05 (آویزان) و ju-09 (دوتایی کافه‌ای) در پوشه عکس
    # ندارند؛ تصویر فعلی‌شان دست‌نخورده می‌ماند تا عکسشان برسد.
}


def main(path):
    files = set(os.listdir(IMAGE_DIR))
    unknown = sorted({name for names in MAPPING.values() for name in names} - files)

    if unknown:
        sys.exit("این فایل‌ها در پوشهٔ تصاویر نیستند: " + "، ".join(unknown))

    workbook = load_workbook(path)
    sheet = workbook["محصولات"]
    header = [cell.value for cell in sheet[1]]
    sku_column = header.index("کد محصول") + 1
    image_column = header.index("تصاویر") + 1

    written = 0
    seen = set()

    for row in range(2, sheet.max_row + 1):
        sku = str(sheet.cell(row=row, column=sku_column).value or "").strip()

        if sku not in MAPPING:
            continue

        seen.add(sku)
        sheet.cell(row=row, column=image_column).value = ", ".join(MAPPING[sku])
        written += 1

    missing = sorted(set(MAPPING) - seen)

    if missing:
        sys.exit("این کدها در فایل اکسل نیستند: " + "، ".join(missing))

    workbook.save(path)

    unused = sorted(files - {name for names in MAPPING.values() for name in names})
    print("تصویر %d محصول نوشته شد." % written)

    if unused:
        print("عکس‌های بی‌استفاده در پوشه: " + "، ".join(unused))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else WORKBOOK)
