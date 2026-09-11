# -*- coding: utf-8 -*-
"""ساخت فایل اکسل خوانا و قابل ویرایش از خروجی CSV ووکامرس.

    python3 tools/export-to-xlsx.py ورودی.csv خروجی.xlsx

ساختار فایل بر پایهٔ کاری چیده شده که واقعاً انجام می‌شود: ویرایش قیمت. پس
ستون‌های قیمت و موجودی نزدیک نام محصول‌اند و متن‌های بلند ته جدول رفته‌اند.

ستون‌هایی که در کل کاتالوگ خالی‌اند حذف می‌شوند تا فایل شلوغ نباشد؛ حذفشان
بی‌خطر است چون افزونه هر ستونی را که نبیند دست نمی‌زند.
"""
import csv
import io
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# (سرستون، ستون مبدأ در فایل ووکامرس، پهنا، قفل، توضیح)
# قفل = ستونی که ساختار محصول به آن وابسته است و نباید ویرایش شود.
COLUMNS = [
    ("شناسه",          "شناسه",             8,  True,  "شناسهٔ محصول در وردپرس. کلید اصلی تطبیق — هرگز تغییرش ندهید."),
    ("نوع",            "نوع",               11, True,  "ساده / متغیر / واریاسیون. ساختار محصول است — تغییرش ندهید."),
    ("کد محصول",       "شناسه محصول",       26, False, "کد کالا (SKU). یکتا. اگر خالی باشد می‌توانید پرش کنید."),
    ("نام",            "نام",               42, False, "عنوان محصول در فروشگاه."),
    ("قیمت",           "قیمت عادی",         13, False, "به تومان. محصول «متغیر» قیمت ندارد — قیمت در سطرهای واریاسیونِ زیرش است."),
    ("قیمت حراج",      "قیمت فروش ویژه",    13, False, "خالی یعنی بدون حراج."),
    ("موجودی",         "انبار",             10, False, "تعداد. خالی یعنی مدیریت انبار خاموش است."),
    ("وضعیت موجودی",   "در انبار؟",         14, False, "موجود / ناموجود."),
    ("دسته‌بندی",      "دسته‌بندی‌ها",      24, False, "چند دسته با کاما؛ زیردسته با «>»."),
    ("وضعیت انتشار",   "منتشر شده",         14, False, "منتشر / پیش‌نویس / خصوصی."),
    ("توضیح کوتاه",    "توضیح کوتاه",       30, False, "یکی دو جمله زیر عنوان محصول."),
    ("توضیحات",        "توضیحات",           48, False, "متن کامل صفحهٔ محصول. HTML مجاز است؛ متن ساده هم خودکار به پاراگراف و فهرست تبدیل می‌شود."),
    ("تصاویر",         "تصاویر",            32, False, "اولی تصویر شاخص، بقیه گالری. با کاما جدا کنید."),
    ("مادر",           "مادر",              22, True,  "فقط برای واریاسیون: کد محصول والد — تغییرش ندهید."),
    ("نام صفت",        "نام 1 صفت",         13, True,  "صفت محصول متغیر — تغییرش ساختار واریاسیون‌ها را می‌شکند."),
    ("مقدار صفت",      "مقدار(های) 1 صفت",  20, True,  "مقادیر صفت — تغییرش ندهید."),
]

TYPE_FA = {"simple": "ساده", "variable": "متغیر", "variation": "واریاسیون",
           "grouped": "گروهی", "external": "خارجی"}
STOCK_FA = {"1": "موجود", "0": "ناموجود"}
STATUS_FA = {"1": "منتشر", "0": "پیش‌نویس", "-1": "خصوصی"}

HEAD_FILL = PatternFill("solid", fgColor="1F2A37")
LOCK_HEAD_FILL = PatternFill("solid", fgColor="6B7280")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)

ROW_FILL = {
    "variable": PatternFill("solid", fgColor="FFF6DF"),   # والد
    "variation": PatternFill("solid", fgColor="F7F7F5"),  # فرزند
}
LOCK_FILL = PatternFill("solid", fgColor="F3F4F6")
EDIT_FILL = PatternFill("solid", fgColor="FFFFFF")
PRICE_FILL = PatternFill("solid", fgColor="FEFCE8")

THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def repair(text):
    """حذف دنباله‌های «\\n» تحت‌اللفظی — همتای PCS_Content::repair در PHP.

    همان نتیجه را می‌دهد: فایل اکسل متن ترمیم‌شده را نشان می‌دهد و افزونه هم هنگام
    ورود دوباره همان را می‌سازد، پس رفت‌وبرگشت تغییری ایجاد نمی‌کند.
    """
    if not text or not text.strip():
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\\r\\n|\\n|\\r", "\n", text)
    text = re.sub(r"\\t", "\t", text)
    text = re.sub(r"[ \t]+$", "", text, flags=re.M)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def main(src, dest):
    with io.open(src, encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))

    # ستون‌های خالی در کل کاتالوگ حذف می‌شوند، مگر ستون‌های کلیدی و ستون‌هایی که
    # ممکن است بخواهید بعداً پرشان کنید (حراج و موجودی).
    essential = {"شناسه", "نوع", "کد محصول", "نام", "قیمت", "قیمت حراج", "موجودی", "وضعیت موجودی"}
    columns = [
        c for c in COLUMNS
        if c[0] in essential or any((r.get(c[1]) or "").strip() for r in rows)
    ]
    dropped = [c[0] for c in COLUMNS if c not in columns]

    wb = Workbook()
    ws = wb.active
    ws.title = "محصولات"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False

    for index, (name, _, width, locked, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index, value=name)
        cell.fill = LOCK_HEAD_FILL if locked else HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.row_dimensions[1].height = 34
    # ستون‌های شناسایی ثابت می‌مانند تا هنگام اسکرول افقی معلوم باشد کدام محصول است.
    # مختصات به‌صورت رشته ساخته می‌شود؛ فراخوانی ws.cell اینجا مکان‌نمای نوشتن را
    # جلو می‌برد و یک سطر خالی پیش از داده‌ها جا می‌گذارد.
    freeze_at = [c[0] for c in columns].index("قیمت") + 1
    ws.freeze_panes = f"{get_column_letter(freeze_at)}2"

    numeric = {"قیمت", "قیمت حراج", "موجودی", "شناسه"}
    long_text = {"توضیحات", "توضیح کوتاه", "تصاویر"}

    # گروه‌بندی واریاسیون‌ها زیر والدشان تا بشود جمعشان کرد.
    groups = []
    current = None

    for record in rows:
        kind = (record.get("نوع") or "").strip()
        values = []

        for name, source, _, _, _ in columns:
            value = (record.get(source) or "").strip()

            if name == "نوع":
                value = TYPE_FA.get(value, value)
            elif name == "وضعیت موجودی":
                value = STOCK_FA.get(value, value)
            elif name == "وضعیت انتشار":
                value = STATUS_FA.get(value, value)
            elif name in ("توضیحات", "توضیح کوتاه"):
                value = repair(value)

            values.append(value)

        ws.append(values)
        index = ws.max_row
        ws.row_dimensions[index].height = 20

        if kind == "variable":
            current = index
            groups.append([index, index])
        elif kind == "variation" and groups and current:
            groups[-1][1] = index
        else:
            current = None

        for column, (name, _, _, locked, _) in enumerate(columns, start=1):
            cell = ws.cell(row=index, column=column)
            cell.border = BORDER
            # متن بلند پیچانده نمی‌شود تا ارتفاع سطرها یکنواخت بماند؛ با کلیک روی
            # سلول، متن کامل در نوار فرمول دیده و ویرایش می‌شود.
            cell.alignment = Alignment(
                horizontal="center" if name in numeric or name in ("نوع", "وضعیت موجودی", "وضعیت انتشار") else "right",
                vertical="center",
                wrap_text=False,
            )

            if locked:
                cell.fill = LOCK_FILL
            elif kind in ROW_FILL:
                cell.fill = ROW_FILL[kind]
            elif name in ("قیمت", "قیمت حراج"):
                cell.fill = PRICE_FILL
            else:
                cell.fill = EDIT_FILL

            if name in ("قیمت", "قیمت حراج"):
                cell.font = Font(bold=True, size=11)

            if name in numeric:
                text = str(cell.value or "").strip()
                if text.replace(".", "", 1).isdigit():
                    cell.value = float(text) if "." in text else int(text)
                    if name != "شناسه":
                        cell.number_format = "#,##0"

            if name in long_text:
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)

    # جمع‌شدنی کردن واریاسیون‌ها.
    for start, end in groups:
        if end > start:
            ws.row_dimensions.group(start + 1, end, outline_level=1, hidden=False)

    last_column = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_column}{ws.max_row}"

    # فهرست‌های کشویی برای ستون‌های وضعیت.
    for name, options in (("وضعیت موجودی", "موجود,ناموجود"), ("وضعیت انتشار", "منتشر,پیش‌نویس,خصوصی")):
        if not any(c[0] == name for c in columns):
            continue
        letter = get_column_letter([c[0] for c in columns].index(name) + 1)
        rule = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{ws.max_row}")

    build_guide(wb, columns, dropped, len(rows))
    wb.save(dest)
    print(f"ساخته شد: {dest}  ({len(rows)} سطر، {len(columns)} ستون)")
    if dropped:
        print("ستون‌های خالی حذف‌شده: " + "، ".join(dropped))


def build_guide(wb, columns, dropped, total):
    guide = wb.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    guide.sheet_view.showGridLines = False

    def heading(text):
        guide.append([text, ""])
        guide.cell(row=guide.max_row, column=1).font = Font(bold=True, size=12, color="1F2A37")

    def line(key, value=""):
        guide.append([key, value])
        guide.cell(row=guide.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    heading("چطور کار کنم؟")
    line("۱", "قیمت یا هر مقدار دیگری را در همین فایل عوض کنید و ذخیره کنید.")
    line("۲", "پیشخوان ← محصولات ← همگام‌سازی با اکسل ← بارگذاری فایل.")
    line("۳", "جدول پیش‌نمایش را بخوانید — دقیقاً می‌گوید چه چیزی از چه به چه تغییر می‌کند.")
    line("۴", "اگر درست بود، «اعمال» را بزنید. تا آن لحظه هیچ‌چیز در سایت عوض نمی‌شود.")
    guide.append([])

    heading("قاعده‌های مهم")
    line("سلول خالی", "یعنی «این فیلد را تغییر نده» — مقدار فعلی محصول پاک نمی‌شود.")
    line("ستون‌های خاکستری", "ساختار محصول به آن‌ها وابسته است؛ تغییرشان محصول را خراب می‌کند.")
    line("قیمت‌ها", "به تومان‌اند — همان عددی که در سایت دیده می‌شود.")
    line("محصول «متغیر»", "خودش قیمت ندارد. قیمت واقعی در سطرهای «واریاسیون» زیر آن است.")
    line("جمع کردن واریاسیون‌ها", "با دکمهٔ «−» کنار شمارهٔ سطرها می‌توانید واریاسیون‌ها را ببندید.")
    line("فیلتر", "روی سرستون‌ها فیلتر فعال است؛ مثلاً فقط یک دسته را نشان دهید.")
    guide.append([])

    heading("نوشتن توضیحات")
    line("متن ساده", "خط خالی بین پاراگراف‌ها بگذارید؛ خودکار به پاراگراف تبدیل می‌شود.")
    line("فهرست", "هر خط را با «- » یا «* » شروع کنید تا فهرست نقطه‌ای شود.")
    line("HTML", "اگر تگ بنویسید (مثل <strong> یا <ul>) دست‌نخورده باقی می‌ماند.")
    guide.append([])

    heading("ستون‌ها")
    for name, _, _, locked, note in columns:
        line(("🔒 " if locked else "") + name, note)

    if dropped:
        guide.append([])
        heading("ستون‌های حذف‌شده")
        line("چرا", "این ستون‌ها در کل کاتالوگ خالی بودند و برای خوانایی حذف شدند: "
                   + "، ".join(dropped))
        line("عیبی ندارد؟", "بله — افزونه هر ستونی را که در فایل نباشد دست نمی‌زند.")

    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 86
    for row in range(1, guide.max_row + 1):
        guide.cell(row=row, column=1).alignment = Alignment(vertical="top", horizontal="right")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
